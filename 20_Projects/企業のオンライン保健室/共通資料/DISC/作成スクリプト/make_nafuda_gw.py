from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = r"C:\Users\y-takahashi\MyBrain\20_Projects\企業のオンライン保健室\顧客\有限会社グリットワークス\01_メンバーカルテ\グリットワークス_トリセツ名札_v0.2.pptx"
FONT = "游ゴシック"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x77, 0x77, 0x77)
LIGHT = RGBColor(0xEE, 0xEE, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

TYPES = [
    (0, 0, "切りひらく", RGBColor(0xE8, 0x6A, 0x3A)),
    (0, 1, "組み立てる", RGBColor(0x3A, 0x7C, 0xC4)),
    (1, 0, "支える", RGBColor(0x4E, 0xA8, 0x6B)),
    (1, 1, "つなぐ", RGBColor(0xC2, 0x5B, 0xA8)),
]

# パワーカラー：(帯の色, 帯の上の文字色, 濃い色)
COLORS = {
    "紫": (RGBColor(0x8E, 0x5D, 0xB5), WHITE, RGBColor(0x5E, 0x35, 0x80)),
    "赤": (RGBColor(0xD9, 0x45, 0x3B), WHITE, RGBColor(0x9E, 0x24, 0x1C)),
    "緑": (RGBColor(0x5D, 0xAE, 0x6B), WHITE, RGBColor(0x2F, 0x6E, 0x3B)),
    "藍": (RGBColor(0x2E, 0x4A, 0x8B), WHITE, RGBColor(0x1C, 0x2E, 0x5A)),
    "橙": (RGBColor(0xF2, 0x99, 0x4A), NAVY, RGBColor(0xA8, 0x5A, 0x12)),
}
ENGINE = {"純": "気持ちに正直", "海": "しなやか・柔軟", "樹": "やさしさ",
          "旅": "新しいこと好き", "陽": "まじめ・誠実"}

# (氏名, パワーカラー, 色の説明, 深層心理, アクティビティレベル, シンキングパターン, ひとこと（仮）, ラッキーナンバー)
# 利き脳：(行, 列)。未判定は登録しない
KIKINOU = {"大関 謙吾": (1, 0)}  # 義＝親指右・腕右
PEOPLE = [
    ("林 範世", "紫", "興味心旺盛・旅や移動が好き", "海", 41275248, 61,
     "新しい話、大好きです。\n理由と全体像をもらえると早いです！", 3),
    ("文屋 伴張", "赤", "情熱的・リーダー的", "樹", 483253680, 73,
     "結論と理由から話してください。\n先の見通しが見えると力が出ます！", 7),
    ("小藪 麗奈", "赤", "情熱的・リーダー的", "旅", 194706502, 46,
     "要点を先に、気持ちも一緒に。\n新しいことは任せてください！", 4),
    ("藤井 和子", "緑", "癒し系・物静かでも心は熱い", "純", 16780451, 22,
     "気持ちを分かってもらえると頑張れます。\n考える時間を少しください。", 5),
    ("大関 謙吾", "藍", "お洒落・器用", "陽", 26489140, 42,
     "しっかり準備して臨みたいタイプです。\n気持ちのひと言があるとうれしいです！", 5),
    ("渥美 沙弥香", "橙", "まじめ・サポート役が得意", "樹", 444981600, 38,
     "サポート役、得意です。\nテンポよく、気持ちも添えて伝えてください！", 8),
]

prs = Presentation()
prs.slide_width = Cm(9.1)
prs.slide_height = Cm(5.5)
blank = prs.slide_layouts[6]


def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    s = sl.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def text(sl, x, y, w, h, t, size=7, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = sl.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(t.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color
    return tb


def meter(sl, x, y, w, label, left, right, pos, note, bar, dot):
    text(sl, x, y, w, 0.35, label, size=6, bold=True, color=GRAY)
    by = y + 0.42
    box(sl, x, by, w, 0.14, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box(sl, x, by, max(w * pos, 0.14), 0.14, fill=bar, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    d = 0.3
    box(sl, x + w * pos - d / 2, by - 0.08, d, d, fill=dot, line=WHITE, shape=MSO_SHAPE.OVAL, lw=1)
    text(sl, x, by + 0.2, w / 2, 0.3, left, size=5.5, color=GRAY)
    text(sl, x + w / 2, by + 0.2, w / 2, 0.3, right, size=5.5, color=GRAY, align=PP_ALIGN.RIGHT)
    text(sl, x, by + 0.48, w, 0.3, note, size=6, bold=True, color=NAVY)


def act_pos(v):
    # プロファイリングの区分（0/200M/300M/500M/900M/8,700M）を目盛りにならべる
    pts = [(0, 0), (200e6, 0.2), (300e6, 0.35), (500e6, 0.55), (900e6, 0.75), (8.7e9, 1.0)]
    for (a, pa), (b, pb) in zip(pts, pts[1:]):
        if v <= b:
            return max(0.03, pa + (v - a) / (b - a) * (pb - pa))
    return 1.0


def act_note(v):
    if v < 200e6:
        return "思考型：納得するまで集めてから動く"
    if v < 300e6:
        return "思考＋直感型：動きながら集めて考える"
    if v < 500e6:
        return "直感型：人より少し早く動き出す"
    if v < 900e6:
        return "直感＋体験型：考えより体が先に動く"
    return "体験＋衝動型：まず体が動く"


def tp_note(t):
    if t < 40:
        return "フィーリング：気持ちで受け取る"
    if t < 70:
        return "両方：気持ちも理屈も使える"
    return "ロジカル：理屈・データで受け取る"


for name, pc, pc_desc, engine, act, tp, hitokoto, lucky in PEOPLE:
    band, ink, dark = COLORS[pc]
    s = prs.slides.add_slide(blank)
    box(s, 0, 0, 9.1, 5.5, fill=WHITE)
    box(s, 0, 0, 9.1, 1.55, fill=band)
    text(s, 0.35, 0.12, 5.5, 0.3, "企業のオンライン保健室　はたらくトリセツ", size=5.5, color=ink)
    text(s, 0.35, 0.42, 6.3, 0.75, name, size=17, bold=True, color=ink)
    text(s, 0.35, 1.13, 6.3, 0.3, f"パワーカラー {pc}：{pc_desc}", size=5.5, color=ink)
    box(s, 6.95, 0.22, 1.8, 1.1, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 6.95, 0.25, 1.8, 0.3, "心のエンジン", size=5, color=GRAY, align=PP_ALIGN.CENTER)
    text(s, 6.95, 0.5, 1.8, 0.55, engine, size=15, bold=True, color=dark, align=PP_ALIGN.CENTER)
    text(s, 6.95, 1.02, 1.8, 0.25, ENGINE[engine], size=4.5, color=GRAY, align=PP_ALIGN.CENTER)

    gx, gy, cw, ch = 0.35, 1.95, 1.55, 0.95
    my = KIKINOU.get(name)
    text(s, gx, 1.65, 3.3, 0.3, "利き脳タイプ" if my else "利き脳タイプ（未判定）", size=6, bold=True, color=GRAY)
    for r_, c_, tname, col in TYPES:
        mine = (r_, c_) == my
        x = gx + c_ * (cw + 0.08)
        y = gy + r_ * (ch + 0.08)
        box(s, x, y, cw, ch, fill=col if mine else LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, x, y + 0.08, cw, 0.5, tname, size=7 if mine else 6.5, bold=mine,
             color=WHITE if mine else GRAY, align=PP_ALIGN.CENTER)
        text(s, x, y + 0.55, cw, 0.3, "の人 ★あなた" if mine else "の人", size=5.5, bold=mine,
             color=WHITE if mine else GRAY, align=PP_ALIGN.CENTER)

    mx, mw = 3.85, 4.9
    meter(s, mx, 1.65, mw, "動き出し方（アクティビティレベル）", "じっくり考えて", "まず体が動く",
          act_pos(act), act_note(act), band, dark)
    meter(s, mx, 2.9, mw, f"伝わり方（シンキングパターン {tp}）", "気持ち・感動", "理屈・データ",
          (tp - 11) / 88, tp_note(tp), band, dark)

    box(s, 0.35, 4.2, 8.4, 0.95, fill=RGBColor(0xF4, 0xF4, 0xF4), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0.55, 4.25, 1.6, 0.85, "わたしへの\nひとこと", size=5.5, bold=True, color=dark)
    text(s, 2.0, 4.25, 5.6, 0.85, hitokoto, size=7, bold=True, color=NAVY)
    text(s, 7.5, 4.3, 1.15, 0.75, f"Lucky\n{lucky}", size=6.5, bold=True, color=dark, align=PP_ALIGN.CENTER)

# ================= 裏面（全員共通の見方） =================
s = prs.slides.add_slide(blank)
box(s, 0, 0, 9.1, 5.5, fill=WHITE)
box(s, 0, 0, 9.1, 0.6, fill=NAVY)
text(s, 0.35, 0.05, 8.4, 0.5, "トリセツ名札の見方", size=9, bold=True, color=WHITE)
text(s, 0.35, 0.7, 8.4, 0.35, "利き脳タイプ：指を組んで下になる親指＝受け取り方／腕を組んで下になる腕＝出し方", size=5.5, color=GRAY)
desc = {
    "切りひらく": "親指 左・腕 右｜新しいことの最初の一歩",
    "組み立てる": "親指 左・腕 左｜見通しと段取りで形にする",
    "支える": "親指 右・腕 右｜役割が決まると力を尽くす",
    "つなぐ": "親指 右・腕 左｜みんなで動くと場が明るい",
}
for i, (r_, c_, name, col) in enumerate(TYPES):
    y = 1.1 + i * 0.55
    box(s, 0.35, y, 1.7, 0.45, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0.35, y, 1.7, 0.45, name + "人", size=6.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, 2.2, y, 6.6, 0.45, desc[name], size=6, color=NAVY)
text(s, 0.35, 3.35, 8.4, 0.35, "上の帯の色＝パワーカラー（外から見える印象）／心のエンジン＝潜在意識", size=5.5, color=GRAY)
text(s, 0.35, 3.7, 8.4, 0.35, "メーター＝動き出し方・伝わり方。真ん中が正解ではなく、どの位置もその人の持ち味です", size=5.5, color=GRAY)
box(s, 0.35, 4.2, 8.4, 0.95, fill=RGBColor(0xEA, 0xF2, 0xFA), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.55, 4.25, 8.0, 0.85, "これは評価ではありません。関わり方のヒントです。\n「この人にはどう伝えるといいかな？」と思った時に見てください。", size=6.5, bold=True, color=NAVY)

prs.save(OUT)
print(OUT)
