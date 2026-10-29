from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

OUT = r"C:\Users\y-takahashi\MyBrain\20_Projects\企業のオンライン保健室\共通資料\DISC\見本_トリセツ名札_髙橋由香_v0.1.pptx"
FONT = "游ゴシック"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x77, 0x77, 0x77)
LIGHT = RGBColor(0xEE, 0xEE, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
POWER = RGBColor(0xF5, 0xC5, 0x18)      # パワーカラー黄
POWER_DARK = RGBColor(0x9A, 0x74, 0x00)
FIRE = RGBColor(0xE0, 0x4A, 0x2A)

# 利き脳4タイプ（ゆっちぃ仮名）と色
TYPES = [  # (行, 列, 名前, 色)
    (0, 0, "切りひらく", RGBColor(0xE8, 0x6A, 0x3A)),
    (0, 1, "組み立てる", RGBColor(0x3A, 0x7C, 0xC4)),
    (1, 0, "支える", RGBColor(0x4E, 0xA8, 0x6B)),
    (1, 1, "つなぐ", RGBColor(0xC2, 0x5B, 0xA8)),
]
MY_TYPE = (0, 0)

prs = Presentation()
prs.slide_width = Cm(9.1)
prs.slide_height = Cm(5.5)
blank = prs.slide_layouts[6]


def box(sl, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    s = sl.shapes.add_shape(shape, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def text(sl, x, y, w, h, t, size=7, bold=False, color=NAVY, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = sl.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    lines = t.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = color
    return tb


def meter(sl, x, y, w, label, left, right, pos, note):
    """pos: 0.0〜1.0 の位置に●"""
    text(sl, x, y, w, 0.35, label, size=6, bold=True, color=GRAY)
    by = y + 0.42
    box(sl, x, by, w, 0.14, fill=LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    box(sl, x, by, max(w * pos, 0.14), 0.14, fill=POWER, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    d = 0.3
    box(sl, x + w * pos - d / 2, by - 0.08, d, d, fill=POWER_DARK, line=WHITE, shape=MSO_SHAPE.OVAL, lw=1)
    text(sl, x, by + 0.2, w / 2, 0.3, left, size=5.5, color=GRAY)
    text(sl, x + w / 2, by + 0.2, w / 2, 0.3, right, size=5.5, color=GRAY, align=PP_ALIGN.RIGHT)
    text(sl, x, by + 0.48, w, 0.3, note, size=6, bold=True, color=NAVY)


# ================= 表面 =================
s = prs.slides.add_slide(blank)
box(s, 0, 0, 9.1, 5.5, fill=WHITE)
# 上の帯＝パワーカラー
box(s, 0, 0, 9.1, 1.55, fill=POWER)
text(s, 0.35, 0.12, 5.5, 0.3, "企業のオンライン保健室　はたらくトリセツ", size=5.5, color=NAVY)
text(s, 0.35, 0.42, 5.6, 0.75, "髙橋 由香", size=17, bold=True, color=NAVY)
text(s, 0.35, 1.13, 5.6, 0.3, "パワーカラー 黄：どうしても目立つ・楽しいこと優先", size=5.5, color=NAVY)
# 炎バッジ
box(s, 6.95, 0.22, 1.8, 1.1, fill=WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 6.95, 0.25, 1.8, 0.3, "心のエンジン", size=5, color=GRAY, align=PP_ALIGN.CENTER)
text(s, 6.95, 0.5, 1.8, 0.55, "炎", size=15, bold=True, color=FIRE, align=PP_ALIGN.CENTER)
text(s, 6.95, 1.02, 1.8, 0.25, "情熱・喜怒哀楽が大きい", size=4.5, color=GRAY, align=PP_ALIGN.CENTER)

# 左：利き脳4マス
gx, gy, cw, ch = 0.35, 1.95, 1.55, 0.95
text(s, gx, 1.65, 3.2, 0.3, "利き脳タイプ", size=6, bold=True, color=GRAY)
for r_, c_, name, col in TYPES:
    mine = (r_, c_) == MY_TYPE
    x = gx + c_ * (cw + 0.08)
    y = gy + r_ * (ch + 0.08)
    b = box(s, x, y, cw, ch, fill=col if mine else LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, y + 0.08, cw, 0.5, name, size=7 if mine else 6.5,
         bold=mine, color=WHITE if mine else GRAY, align=PP_ALIGN.CENTER)
    text(s, x, y + 0.55, cw, 0.3, "の人 ★あなた" if mine else "の人", size=5.5, bold=mine, color=WHITE if mine else GRAY, align=PP_ALIGN.CENTER)

# 右：メーター
mx, mw = 3.85, 4.9
meter(s, mx, 1.65, mw, "動き出し方（アクティビティレベル）", "じっくり考えて", "まず体が動く", 0.04, "思考型：納得するまで集めてから動く")
meter(s, mx, 2.9, mw, "伝わり方（シンキングパターン 21）", "気持ち・感動", "理屈・データ", (21 - 11) / 88, "フィーリング：気持ちで受け取る")

# 下：ひとこと
box(s, 0.35, 4.2, 8.4, 0.95, fill=RGBColor(0xFF, 0xF6, 0xD6), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.55, 4.25, 1.6, 0.85, "わたしへの\nひとこと", size=5.5, bold=True, color=POWER_DARK)
text(s, 2.0, 4.25, 5.6, 0.85, "材料をたっぷりちょうだい。\n心が決まったら、先頭で走ります！", size=7, bold=True, color=NAVY)
text(s, 7.5, 4.3, 1.15, 0.75, "Lucky\n3", size=6.5, bold=True, color=POWER_DARK, align=PP_ALIGN.CENTER)

# ================= 裏面（見方） =================
s = prs.slides.add_slide(blank)
box(s, 0, 0, 9.1, 5.5, fill=WHITE)
box(s, 0, 0, 9.1, 0.6, fill=NAVY)
text(s, 0.35, 0.05, 8.4, 0.5, "トリセツ名札の見方", size=9, bold=True, color=WHITE)
text(s, 0.35, 0.7, 8.4, 0.35, "利き脳タイプ：指を組んで下になる親指＝受け取り方／腕を組んで下になる腕＝出し方", size=5.5, color=GRAY)
desc = {
    "切りひらく": "親指 左・腕 右｜新しいことの最初の一歩",
    "組み立てる": "親指 左・腕 左｜見通しと段取りで形にする",
    "支える": "親指 右・腕 右｜役割が決まると力を尽くす",
    "つなぐ": "親指 右・腕 左｜みんなで動くと場が明るい",
}
for i, (r_, c_, name, col) in enumerate(TYPES):
    y = 1.1 + i * 0.55
    box(s, 0.35, y, 1.7, 0.45, fill=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 0.35, y, 1.7, 0.45, name + "人", size=6.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    text(s, 2.2, y, 6.6, 0.45, desc[name], size=6, color=NAVY)
text(s, 0.35, 3.35, 8.4, 0.35, "上の帯の色＝パワーカラー（外から見える印象）／心のエンジン＝潜在意識", size=5.5, color=GRAY)
text(s, 0.35, 3.7, 8.4, 0.35, "メーター＝動き出し方・伝わり方。真ん中が正解ではなく、どの位置もその人の持ち味です", size=5.5, color=GRAY)
box(s, 0.35, 4.2, 8.4, 0.95, fill=RGBColor(0xEA, 0xF2, 0xFA), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 0.55, 4.25, 8.0, 0.85, "これは評価ではありません。関わり方のヒントです。\n「この人にはどう伝えるといいかな？」と思った時に見てください。", size=6.5, bold=True, color=NAVY)

prs.save(OUT)
print(OUT)
