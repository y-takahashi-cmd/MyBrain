import sys
from pptx import Presentation
src,dst=sys.argv[1],sys.argv[2]
R=[('仁部 あきえ さん　90分','〇〇 さん　90〜120分'),
   ('仁部さんは話すこと','ご本人は話すこと'),('仁部さんは話すだけ','ご本人は話すだけ'),
   ('お父さんのお酒・家にお金を入れない・毎日の大喧嘩の話が出たら、ワークの場面候補としてメモ。','お金にまつわる子どもの頃の場面が出たら、ワークの場面候補としてメモ。'),
   ('「お金がない」','「＿＿＿＿＿＿」'),('「家にお金を入れない」','（事前に聞いたエピソードを入れる）'),
   ('仁部さんが話してくれたエピソード。','事前に聞いたご本人のエピソードを上の2行に入れる。'),
   ('いま聴いた話（お父さんのお酒・両親の喧嘩など）','いま聴いた子どもの頃の話'),
   ('第1部で聴いた仁部さんの話','第1部で聴いたご本人の話'),
   ('第1部で聴いたお父さん・お母さんの話','第1部で聴いた子どもの頃の話'),
   ('お父さんのお酒・両親の喧嘩の場面が出る可能性が高い。','第1部で聴いた子どもの頃の場面が出やすい。')]
p=Presentation(src); hit={a:0 for a,_ in R}
def fix(tf):
    for para in tf.paragraphs:
        for r in para.runs:
            for a,b in R:
                if a in r.text: r.text=r.text.replace(a,b); hit[a]+=1
for s in p.slides:
    for sh in s.shapes:
        if sh.has_text_frame: fix(sh.text_frame)
    if s.has_notes_slide: fix(s.notes_slide.notes_text_frame)
p.save(dst); print({k[:12]:v for k,v in hit.items()})
